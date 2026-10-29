"""
rerun_missing_answers.py
========================
Fixes rows in the answers file where the system answer is [ERROR: 429 ...]
(API credits ran out mid-benchmark), then re-judges those rows.

Dependency tiers
----------------
  System A : only openai + openpyxl            (runs on IIMB laptop / anywhere)
  System B : openai + chromadb + HF embeddings (needs torch; blocked by WDAC on IIMB)
  System C : openai + chromadb + HF embeddings (same)

The script auto-detects which imports are available and regenerates whatever it can.
Missing imports produce a warning, not a crash.

Usage
-----
  # From 05_App/benchmark/  (paths auto-detected)
  python rerun_missing_answers.py

  # Explicit paths
  python rerun_missing_answers.py \
      --answers  results/benchmark_answers_20261024_185733.xlsx \
      --eval     results/benchmark_evaluation_20261029_032200.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

# env loading
HERE = Path(__file__).resolve().parent

def _load_env():
    try:
        from dotenv import load_dotenv
        load_dotenv(HERE / ".env")
        load_dotenv(HERE.parent / ".env", override=True)
    except ImportError:
        pass

_load_env()

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MODEL = os.getenv("LLM_MODEL", "gpt-5-nano")


def _make_openai_client():
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        sys.exit("ERROR: OPENAI_API_KEY not set. Check 05_App/.env.")
    from openai import OpenAI
    return OpenAI(api_key=api_key)


def _load_chromadb_store():
    try:
        from langchain_huggingface import HuggingFaceEmbeddings
        from langchain_chroma import Chroma
        index_dir = HERE.parent / "chroma_db"
        if not index_dir.exists():
            print("  [WARN] chroma_db/ not found -- skipping Systems B & C.")
            return None
        emb = HuggingFaceEmbeddings(
            model_name="BAAI/bge-small-en-v1.5",
            encode_kwargs={"normalize_embeddings": True},
        )
        store = Chroma(
            collection_name="regulatory_corpus",
            embedding_function=emb,
            persist_directory=str(index_dir),
        )
        print(f"  ChromaDB loaded: {index_dir}")
        return store
    except Exception as exc:
        print(f"  [WARN] ChromaDB unavailable ({exc}).")
        print("         Systems B & C will be skipped on this machine.")
        return None


SYSTEM_A_PROMPT = """\
You are an Indian regulatory compliance expert. Answer the question using your
general knowledge of Indian laws and regulations. Be specific, cite relevant
sections where known, and flag any uncertainty. Structure your answer as:
1. Direct answer
2. Relevant legal provision (Act / Section / Rule)
3. Key compliance takeaway
"""

def _gen_system_a(client, question_text):
    rsp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": SYSTEM_A_PROMPT},
            {"role": "user",   "content": question_text},
        ],
    )
    return rsp.choices[0].message.content.strip()


SYSTEM_B_PROMPT = """\
You are an Indian regulatory compliance assistant. Answer ONLY from the
retrieved sources supplied. Structure every answer as:
1. A direct answer.
2. Sources relied on, cited by document TITLE.
3. The specific regulatory provision (section, rule, clause).
4. A concise compliance takeaway.
Hard rules: never invent citations; flag DRAFT material; if evidence is
insufficient, say so explicitly."""

def _gen_system_b(client, question_text, domain, store):
    TOP_K, CANDIDATE_K = 6, 24
    TIER_WEIGHT = {"A": 1.00, "B": 0.95, "C": 0.80, "D": 0.55, "E": 0.35}
    results = store.similarity_search_with_relevance_scores(
        question_text, k=CANDIDATE_K, filter={"domain": domain}
    )
    scored = []
    for doc, raw in results:
        md = doc.metadata or {}
        if md.get("tier", "E") == "E":
            continue
        space = (store._collection.metadata or {}).get("hnsw:space", "l2").lower()
        cos = raw if space.startswith("cosine") else 1.0 - (1.0 - raw) ** 2
        scored.append({
            "content": doc.page_content,
            "title": md.get("title", "Unknown"),
            "tier": md.get("tier", "?"),
            "weighted": cos * TIER_WEIGHT.get(md.get("tier", "C"), 0.80),
        })
    scored.sort(key=lambda x: x["weighted"], reverse=True)
    context = "\n\n".join(
        f"[{i}] Title: {p['title']} | Tier: {p['tier']}\n{p['content']}"
        for i, p in enumerate(scored[:TOP_K], 1)
    )
    rsp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": SYSTEM_B_PROMPT},
            {"role": "user", "content": f"Domain: {domain}\n\nRetrieved sources:\n{context}\n\nQuestion: {question_text}"},
        ],
    )
    return rsp.choices[0].message.content.strip()


def _gen_system_c(client, question_text, domain, store):
    try:
        sys.path.insert(0, str(HERE.parent))
        from agentic_rag import agentic_answer
        result = agentic_answer(question_text, domain, store=store, client=client)
        return result.get("answer", "")
    except Exception as exc:
        return f"[ERROR: {exc}]"


JUDGE_SYSTEM = """\
You are an expert evaluator for an Indian regulatory compliance Q&A system.
Score the SYSTEM ANSWER against the EXPECTED ANSWER and the QUESTION on three
dimensions using the rubric below. Be strict and objective.

RUBRIC
------
Factual Accuracy (0-4)
  4 = All key facts correct, no hallucinations, sections/penalties cited correctly.
  3 = Mostly correct with minor factual gaps or small inaccuracies.
  2 = Partially correct; some material facts wrong or missing.
  1 = Mostly incorrect; fundamental misunderstanding of the regulation.
  0 = Completely wrong or no relevant content.

Completeness (0-2)
  2 = Covers all key aspects: legal provision, compliance requirement, practical impact.
  1 = Covers some aspects but misses important ones.
  0 = Very incomplete or single-dimensional.

Confidence Calibration (0-2)
  2 = Uncertainty expressed appropriately; no overconfidence where knowledge is thin;
      hedges on draft/proposed rules; does not claim certainty on unclear areas.
  1 = Partially calibrated -- some overconfidence or underconfidence.
  0 = Significantly overconfident (states wrong facts as certain) or so hedged
      as to be useless.

Return ONLY valid JSON -- no markdown, no preamble:
{
  "factual_accuracy":        <integer 0-4>,
  "completeness":            <integer 0-2>,
  "confidence_calibration":  <integer 0-2>,
  "total":                   <integer 0-8>,
  "rationale":               "<1-2 sentences>"
}
"""

def _judge(client, question, expected, answer):
    if not answer or answer.startswith("[ERROR"):
        return {
            "factual_accuracy": 0, "completeness": 0,
            "confidence_calibration": 0, "total": 0,
            "rationale": "No answer produced (error or skipped).",
        }
    user_msg = (
        f"QUESTION:\n{question}\n\n"
        f"EXPECTED ANSWER (ground truth):\n{expected}\n\n"
        f"SYSTEM ANSWER (to evaluate):\n{answer}"
    )
    try:
        stream = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": JUDGE_SYSTEM},
                {"role": "user",   "content": user_msg},
            ],
            stream=True,
        )
        raw = ""
        for chunk in stream:
            if chunk.choices and chunk.choices[0].delta.content:
                raw += chunk.choices[0].delta.content
        raw = re.sub(r"^```(?:json)?\s*", "", raw.strip())
        raw = re.sub(r"\s*```$", "", raw).strip()
        result = json.loads(raw)
        result["total"] = (
            result.get("factual_accuracy", 0)
            + result.get("completeness", 0)
            + result.get("confidence_calibration", 0)
        )
        return result
    except Exception as exc:
        return {
            "factual_accuracy": 0, "completeness": 0,
            "confidence_calibration": 0, "total": 0,
            "rationale": f"Judge error: {exc}",
        }


def _load_sheet(path):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    return headers, rows


def _save_answers(headers, rows, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    col_widths = {
        "Question Id": 14, "Domain": 12, "Tier": 6,
        "Question Text": 55, "Expected Answer": 70,
        "System A Answer": 70, "System B Answer": 70, "System C Answer": 70,
    }
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(1, ci, h)
        cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(h, 14)
    ws.row_dimensions[1].height = 25
    domain_fills = {
        "Healthcare": PatternFill("solid", fgColor="E2EFDA"),
        "RERA":       PatternFill("solid", fgColor="DDEEFF"),
    }
    for ri, row in enumerate(rows, 2):
        fill = domain_fills.get(str(row.get("Domain", "") or ""), PatternFill())
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(ri, ci, row.get(h, ""))
            cell.font = Font(name="Calibri", size=9)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        ws.row_dimensions[ri].height = 80
    ws.freeze_panes = "A2"
    wb.save(str(path))
    print(f"Saved answers: {path}")


def _save_eval(headers, rows, path):
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    sys_fills = {
        "A": PatternFill("solid", fgColor="FFF2CC"),
        "B": PatternFill("solid", fgColor="DDEEFF"),
        "C": PatternFill("solid", fgColor="E2EFDA"),
    }
    col_widths = {
        "Question Id": 14, "Domain": 12, "Tier": 6, "System": 8,
        "Factual Accuracy": 10, "Completeness": 10, "Confidence Calibration": 14,
        "Total": 8, "Rationale": 60, "Question Text": 50,
        "Expected Answer": 55, "System Answer": 55,
    }
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scores"
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(1, ci, h)
        cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(h, 14)
    ws.row_dimensions[1].height = 28
    for ri, row in enumerate(rows, 2):
        fill = sys_fills.get(str(row.get("System", "") or ""), PatternFill())
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(ri, ci, row.get(h, ""))
            cell.font = Font(name="Calibri", size=9)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        ws.row_dimensions[ri].height = 60
    ws.freeze_panes = "A2"
    wb.save(str(path))
    print(f"Saved eval: {path}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--answers", type=Path,
                        default=HERE / "results" / "benchmark_answers_20261024_185733.xlsx")
    parser.add_argument("--eval", type=Path,
                        default=HERE / "results" / "benchmark_evaluation_20261029_032200.xlsx")
    args = parser.parse_args()

    if not args.answers.exists():
        sys.exit(f"ERROR: {args.answers} not found")
    if not args.eval.exists():
        sys.exit(f"ERROR: {args.eval} not found")

    client = _make_openai_client()
    print(f"OpenAI client ready. Model: {MODEL}")

    print("\nLoading ChromaDB (needed for Systems B & C) ...")
    store = _load_chromadb_store()

    # Load answers
    print(f"\nLoading answers: {args.answers}")
    ans_headers, ans_rows = _load_sheet(args.answers)
    ans_by_qid = {str(r.get("Question Id") or ""): r for r in ans_rows}

    # Find which cells need regeneration
    need_regen = []
    for r in ans_rows:
        qid = str(r.get("Question Id") or "")
        for col, label in [("System A Answer","A"), ("System B Answer","B"), ("System C Answer","C")]:
            if str(r.get(col, "") or "").startswith("[ERROR"):
                need_regen.append((qid, label))

    can_do  = [(q,s) for q,s in need_regen if s == "A" or store is not None]
    cant_do = [(q,s) for q,s in need_regen if s != "A" and store is None]

    print(f"\nTotal error cells  : {len(need_regen)}")
    print(f"  Will regenerate  : {len(can_do)}")
    print(f"  Skipping (no ChromaDB): {len(cant_do)}")

    if not can_do:
        print("\nNothing to regenerate. Exiting.")
        return

    # Step 1: regenerate answers
    print(f"\n{'='*60}")
    print("STEP 1: Regenerating answers")
    print('='*60)
    fixed = 0
    for n, (qid, sys_label) in enumerate(can_do, 1):
        row = ans_by_qid[qid]
        domain = str(row.get("Domain") or "")
        q_text = str(row.get("Question Text") or "")
        col_key = f"System {sys_label} Answer"

        print(f"[{n:3}/{len(can_do)}] {qid} Sys {sys_label} ...", end=" ", flush=True)
        t0 = time.time()
        try:
            if sys_label == "A":
                ans = _gen_system_a(client, q_text)
            elif sys_label == "B":
                ans = _gen_system_b(client, q_text, domain, store)
            else:
                ans = _gen_system_c(client, q_text, domain, store)

            if ans.startswith("[ERROR"):
                print(f"STILL ERROR: {ans[:60]}")
            else:
                ans_by_qid[qid][col_key] = ans
                fixed += 1
                print(f"OK ({time.time()-t0:.1f}s) {len(ans)} chars")
        except Exception as exc:
            print(f"FAILED: {exc}")
        time.sleep(1)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    new_ans_path = args.answers.parent / f"benchmark_answers_{ts}.xlsx"
    _save_answers(ans_headers, list(ans_by_qid.values()), new_ans_path)
    print(f"Fixed {fixed} answer cells.")

    # Step 2: re-judge
    print(f"\n{'='*60}")
    print("STEP 2: Re-judging fixed rows")
    print('='*60)

    eval_headers, eval_rows = _load_sheet(args.eval)

    fixed_pairs = set()
    for qid, sys_label in can_do:
        col = f"System {sys_label} Answer"
        val = str(ans_by_qid[qid].get(col, "") or "")
        if val and not val.startswith("[ERROR"):
            fixed_pairs.add((qid, sys_label))

    rejudge_idx = [
        i for i, row in enumerate(eval_rows)
        if (str(row.get("Question Id") or ""), str(row.get("System") or "")) in fixed_pairs
    ]
    print(f"Rows to re-judge: {len(rejudge_idx)}")

    judged = 0
    for n, idx in enumerate(rejudge_idx, 1):
        row = eval_rows[idx]
        qid = str(row.get("Question Id") or "")
        sys_label = str(row.get("System") or "")
        q_text = str(row.get("Question Text") or "")
        expected = str(row.get("Expected Answer") or "")
        new_ans = str(ans_by_qid[qid].get(f"System {sys_label} Answer", "") or "")

        print(f"[{n:3}/{len(rejudge_idx)}] {qid} Sys {sys_label} ...", end=" ", flush=True)
        t0 = time.time()
        scores = _judge(client, q_text, expected, new_ans)
        print(f"total={scores['total']}/8  ({time.time()-t0:.1f}s)")
        time.sleep(1)

        eval_rows[idx]["System Answer"]          = new_ans
        eval_rows[idx]["Factual Accuracy"]       = scores.get("factual_accuracy", 0)
        eval_rows[idx]["Completeness"]           = scores.get("completeness", 0)
        eval_rows[idx]["Confidence Calibration"] = scores.get("confidence_calibration", 0)
        eval_rows[idx]["Total"]                  = scores.get("total", 0)
        eval_rows[idx]["Rationale"]              = scores.get("rationale", "")
        judged += 1

    final_eval_path = args.eval.parent / f"benchmark_evaluation_{ts}.xlsx"
    _save_eval(eval_headers, eval_rows, final_eval_path)

    still_zero = sum(1 for r in eval_rows if (r.get("Total") or 0) == 0)
    print(f"\n{'='*60}")
    print("DONE")
    print(f"  Answers fixed    : {fixed}")
    print(f"  Rows re-judged   : {judged}")
    print(f"  Rows still 0     : {still_zero}")
    print(f"  New answers file : {new_ans_path}")
    print(f"  New eval file    : {final_eval_path}")
    if cant_do:
        from collections import Counter
        cnt = Counter(s for _, s in cant_do)
        print(f"\n  NOTE: {len(cant_do)} B/C cells skipped (ChromaDB/torch unavailable): {dict(cnt)}")


if __name__ == "__main__":
    main()
