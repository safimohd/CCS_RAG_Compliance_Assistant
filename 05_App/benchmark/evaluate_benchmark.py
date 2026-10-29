"""
Benchmark evaluator — LLM-as-judge scoring for all three systems.

Rubric (8 points total, Prof's criteria)
-----------------------------------------
  Factual Accuracy      0-4   (core facts correct, no hallucination)
  Completeness          0-2   (all key aspects covered)
  Confidence Calibration 0-2  (appropriate uncertainty, no overconfidence)
  ─────────────────────────
  Total                 0-8

For each (question, system) pair the judge returns a JSON dict:
  {
    "factual_accuracy":        <0-4>,
    "completeness":            <0-2>,
    "confidence_calibration":  <0-2>,
    "total":                   <0-8>,
    "rationale":               "<brief explanation>"
  }

Input
-----
  The raw answers file produced by run_benchmark.py:
    benchmark/results/benchmark_answers_*.xlsx

Output
------
  benchmark/results/benchmark_evaluation_{timestamp}.xlsx
    Sheet "Scores"   — one row per (question_id, system) with all sub-scores
    Sheet "Summary"  — per-domain, per-tier, and overall averages for each system
    Sheet "Comparison" — head-to-head comparison table

Usage
-----
  python evaluate_benchmark.py benchmark/results/benchmark_answers_20240101_120000.xlsx
  python evaluate_benchmark.py --help

OPENAI_API_KEY must be set in the environment.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

# Load .env file from the same folder as this script (if it exists)
try:
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).parent / ".env")
except ImportError:
    pass  # dotenv not installed — rely on env vars set manually

HERE    = Path(__file__).resolve().parent
APP_DIR = HERE.parent
sys.path.insert(0, str(APP_DIR))

RESULTS_DIR = HERE / "results"
RESULTS_DIR.mkdir(exist_ok=True)

LLM_MODEL = os.getenv("LLM_MODEL", "gpt-5-nano")

# ─── Judge prompt ──────────────────────────────────────────────────────────────
JUDGE_SYSTEM = """\
You are an expert evaluator for an Indian regulatory compliance Q&A system.
Score the SYSTEM ANSWER against the EXPECTED ANSWER and the QUESTION on three
dimensions using the rubric below. Be strict and objective.

RUBRIC
------
Factual Accuracy (0-4)
  4 = All key facts correct, no hallucinations, sections/penalties cited correctly.
  3 = Mostly correct with minor factual gaps or small inaccuracies.
  2 = Partially correct; some material facts wrong or missing.
  1 = Mostly incorrect; fundamental misunderstanding of the regulation.
  0 = Completely wrong or no relevant content.

Completeness (0-2)
  2 = Covers all key aspects: legal provision, compliance requirement, practical impact.
  1 = Covers some aspects but misses important ones.
  0 = Very incomplete or single-dimensional.

Confidence Calibration (0-2)
  2 = Uncertainty expressed appropriately; no overconfidence where knowledge is thin;
      hedges on draft/proposed rules; does not claim certainty on unclear areas.
  1 = Partially calibrated — some overconfidence or underconfidence.
  0 = Significantly overconfident (states wrong facts as certain) or so hedged
      as to be useless.

Return ONLY valid JSON — no markdown, no preamble:
{
  "factual_accuracy":        <integer 0-4>,
  "completeness":            <integer 0-2>,
  "confidence_calibration":  <integer 0-2>,
  "total":                   <integer 0-8>,
  "rationale":               "<1-2 sentences>"
}
"""


def _judge_one(client, question: str, expected: str, system_answer: str) -> dict:
    """Call LLM judge for one (question, answer) pair."""
    if not system_answer or system_answer.startswith("[ERROR"):
        return {
            "factual_accuracy": 0, "completeness": 0,
            "confidence_calibration": 0, "total": 0,
            "rationale": "No answer produced (error or skipped).",
        }

    user_msg = (
        f"QUESTION:\n{question}\n\n"
        f"EXPECTED ANSWER (ground truth):\n{expected}\n\n"
        f"SYSTEM ANSWER (to evaluate):\n{system_answer}"
    )
    try:
        stream = client.chat.completions.create(
            model=LLM_MODEL,
            messages=[
                {"role": "system", "content": JUDGE_SYSTEM},
                {"role": "user",   "content": user_msg},
            ],
            stream=True,
        )
        raw = ""
        for chunk in stream:
            if chunk.choices and chunk.choices[0].delta.content:
                raw += chunk.choices[0].delta.content
        raw = raw.strip()
        # strip any markdown code fences the model adds despite being told not to
        import re as _re
        raw = _re.sub(r"^```(?:json)?\s*", "", raw)
        raw = _re.sub(r"\s*```$", "", raw).strip()
        result = json.loads(raw)
        # enforce total
        result["total"] = (
            result.get("factual_accuracy", 0)
            + result.get("completeness", 0)
            + result.get("confidence_calibration", 0)
        )
        return result
    except (json.JSONDecodeError, Exception) as exc:
        return {
            "factual_accuracy": 0, "completeness": 0,
            "confidence_calibration": 0, "total": 0,
            "rationale": f"Judge error: {exc}",
        }


def _load_answers(path: Path) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip().lower().replace(" ", "_")
               for c in range(1, ws.max_column + 1)]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rec = {h: (v or "") for h, v in zip(headers, r)}
        if rec.get("question_id"):
            rows.append(rec)
    return rows


def _save_evaluation(score_rows: list[dict], out_path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict

    wb = openpyxl.Workbook()

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="1F4E79")

    # ── Sheet 1: Scores ──────────────────────────────────────────────────────
    ws_scores = wb.active
    ws_scores.title = "Scores"

    score_cols = [
        ("question_id",            14),
        ("domain",                 12),
        ("tier",                    6),
        ("system",                  8),
        ("factual_accuracy",       10),
        ("completeness",           10),
        ("confidence_calibration", 14),
        ("total",                   8),
        ("rationale",              60),
        ("question_text",          50),
        ("expected_answer",        55),
        ("system_answer",          55),
    ]

    for ci, (col, w) in enumerate(score_cols, 1):
        cell = ws_scores.cell(1, ci, col.replace("_", " ").title())
        cell.font      = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = border
        ws_scores.column_dimensions[get_column_letter(ci)].width = w
    ws_scores.row_dimensions[1].height = 28

    sys_fills = {
        "A": PatternFill("solid", fgColor="FFF2CC"),   # yellow
        "B": PatternFill("solid", fgColor="DDEEFF"),   # blue
        "C": PatternFill("solid", fgColor="E2EFDA"),   # green
    }

    for ri, row in enumerate(score_rows, 2):
        fill = sys_fills.get(row.get("system", ""), PatternFill())
        for ci, (col, _) in enumerate(score_cols, 1):
            val = row.get(col, "")
            cell = ws_scores.cell(ri, ci, val)
            cell.font      = Font(name="Calibri", size=9)
            cell.fill      = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border    = border
        ws_scores.row_dimensions[ri].height = 60

    ws_scores.freeze_panes = "A2"

    # ── Sheet 2: Summary ─────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")

    def _avg(vals):
        vals = [v for v in vals if isinstance(v, (int, float))]
        return round(sum(vals) / len(vals), 2) if vals else 0.0

    def _write_header(ws, row, col, text):
        cell = ws.cell(row, col, text)
        cell.font = Font(bold=True, name="Calibri", size=10, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    # Group by (domain, tier, system)
    groups: dict = defaultdict(list)
    for row in score_rows:
        key = (row.get("domain", ""), str(row.get("tier", "")), row.get("system", ""))
        groups[key].append(row)

    r = 1
    ws_sum.cell(r, 1, "CCS Benchmark Evaluation Summary").font = Font(bold=True, size=14)
    r += 2

    for col_i, h in enumerate(
        ["Domain", "Tier", "System", "N", "Avg Factual (0-4)",
         "Avg Complete (0-2)", "Avg Confidence (0-2)", "Avg Total (0-8)"], 1
    ):
        _write_header(ws_sum, r, col_i, h)
    ws_sum.column_dimensions["A"].width = 14
    ws_sum.column_dimensions["B"].width = 8
    ws_sum.column_dimensions["C"].width = 10
    for i in range(4, 9):
        ws_sum.column_dimensions[get_column_letter(i)].width = 20
    r += 1

    # Collect all unique (domain, tier) combos
    all_domains  = sorted({k[0] for k in groups})
    all_tiers    = sorted({k[1] for k in groups}, key=lambda x: int(x) if x.isdigit() else 99)
    all_systems  = ["A", "B", "C"]

    for domain in all_domains:
        for tier in all_tiers:
            for sys_label in all_systems:
                key = (domain, tier, sys_label)
                rows_k = groups.get(key, [])
                if not rows_k:
                    continue
                ws_sum.cell(r, 1, domain)
                ws_sum.cell(r, 2, f"T{tier}")
                ws_sum.cell(r, 3, f"System {sys_label}")
                ws_sum.cell(r, 4, len(rows_k))
                ws_sum.cell(r, 5, _avg([x.get("factual_accuracy", 0) for x in rows_k]))
                ws_sum.cell(r, 6, _avg([x.get("completeness", 0) for x in rows_k]))
                ws_sum.cell(r, 7, _avg([x.get("confidence_calibration", 0) for x in rows_k]))
                ws_sum.cell(r, 8, _avg([x.get("total", 0) for x in rows_k]))
                fill = sys_fills.get(sys_label, PatternFill())
                for ci in range(1, 9):
                    c = ws_sum.cell(r, ci)
                    c.fill = fill
                    c.font = Font(name="Calibri", size=9)
                    c.border = border
                r += 1

    r += 1
    ws_sum.cell(r, 1, "Overall (all tiers)").font = Font(bold=True, name="Calibri")
    r += 1
    for col_i, h in enumerate(
        ["Domain", "System", "N", "Avg Factual", "Avg Complete",
         "Avg Confidence", "Avg Total"], 1
    ):
        _write_header(ws_sum, r, col_i, h)
    r += 1
    for domain in all_domains:
        for sys_label in all_systems:
            domain_rows = [row for row in score_rows
                           if row.get("domain") == domain and row.get("system") == sys_label]
            if not domain_rows:
                continue
            ws_sum.cell(r, 1, domain)
            ws_sum.cell(r, 2, f"System {sys_label}")
            ws_sum.cell(r, 3, len(domain_rows))
            ws_sum.cell(r, 4, _avg([x.get("factual_accuracy", 0) for x in domain_rows]))
            ws_sum.cell(r, 5, _avg([x.get("completeness", 0) for x in domain_rows]))
            ws_sum.cell(r, 6, _avg([x.get("confidence_calibration", 0) for x in domain_rows]))
            ws_sum.cell(r, 7, _avg([x.get("total", 0) for x in domain_rows]))
            fill = sys_fills.get(sys_label, PatternFill())
            for ci in range(1, 8):
                c = ws_sum.cell(r, ci)
                c.fill = fill
                c.font = Font(name="Calibri", size=9)
                c.border = border
            r += 1

    # ── Sheet 3: Comparison ──────────────────────────────────────────────────
    ws_cmp = wb.create_sheet("Comparison")
    ws_cmp.cell(1, 1, "Head-to-Head Comparison (Avg Total Score out of 8)").font = Font(bold=True, size=12)

    # Per domain, per tier matrix
    r = 3
    headers_cmp = ["Domain", "Tier"] + [f"System {s}" for s in all_systems] + ["Best System"]
    for ci, h in enumerate(headers_cmp, 1):
        _write_header(ws_cmp, r, ci, h)
        ws_cmp.column_dimensions[get_column_letter(ci)].width = 14
    r += 1

    from collections import defaultdict as dd
    for domain in all_domains:
        for tier in all_tiers:
            row_scores = {}
            for sys_label in all_systems:
                key = (domain, tier, sys_label)
                rows_k = groups.get(key, [])
                row_scores[sys_label] = _avg([x.get("total", 0) for x in rows_k]) if rows_k else None

            if all(v is None for v in row_scores.values()):
                continue

            ws_cmp.cell(r, 1, domain)
            ws_cmp.cell(r, 2, f"T{tier}")
            best_sys, best_score = None, -1
            for ci, sys_label in enumerate(all_systems, 3):
                val = row_scores[sys_label]
                ws_cmp.cell(r, ci, val if val is not None else "N/A")
                if val is not None and val > best_score:
                    best_score, best_sys = val, sys_label
            ws_cmp.cell(r, len(all_systems) + 3, f"System {best_sys}" if best_sys else "")
            for ci in range(1, len(headers_cmp) + 1):
                c = ws_cmp.cell(r, ci)
                c.font = Font(name="Calibri", size=9)
                c.border = border
            r += 1

    wb.save(str(out_path))


# ─── main ─────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="CCS Benchmark Evaluator")
    parser.add_argument("answers_file", type=Path,
                        help="Path to benchmark_answers_*.xlsx produced by run_benchmark.py")
    parser.add_argument("--domain",  choices=["Healthcare", "RERA"],
                        help="Evaluate only one domain")
    parser.add_argument("--tier",    type=int, action="append", dest="tiers",
                        help="Filter to tier(s) (repeatable)")
    parser.add_argument("--limit",   type=int,
                        help="Max questions to evaluate (for testing)")
    parser.add_argument("--system",  choices=["A", "B", "C"], action="append",
                        dest="systems", help="Evaluate specific system(s) only")
    args = parser.parse_args()

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        sys.exit("ERROR: OPENAI_API_KEY not set in environment.")

    from openai import OpenAI
    llm_base_url = os.getenv("LLM_BASE_URL")
    client = OpenAI(api_key=api_key, base_url=llm_base_url) if llm_base_url else OpenAI(api_key=api_key)

    if not args.answers_file.exists():
        sys.exit(f"ERROR: File not found: {args.answers_file}")

    print(f"Loading answers from: {args.answers_file}")
    answer_rows = _load_answers(args.answers_file)
    print(f"  {len(answer_rows)} answer rows loaded.")

    # Filter
    if args.domain:
        answer_rows = [r for r in answer_rows if r.get("domain") == args.domain]
    if args.tiers:
        answer_rows = [r for r in answer_rows
                       if str(r.get("tier", "")) in {str(t) for t in args.tiers}]
    if args.limit:
        answer_rows = answer_rows[:args.limit]

    systems_to_eval = args.systems or ["A", "B", "C"]
    sys_answer_keys = {
        "A": "system_a_answer",
        "B": "system_b_answer",
        "C": "system_c_answer",
    }

    score_rows: list[dict] = []
    total_calls = len(answer_rows) * len(systems_to_eval)
    call_n = 0

    for row in answer_rows:
        qid      = row.get("question_id", "")
        domain   = row.get("domain", "")
        tier     = row.get("tier", "")
        q_text   = row.get("question_text", "")
        expected = row.get("expected_answer", "")

        for sys_label in systems_to_eval:
            call_n += 1
            answer = str(row.get(sys_answer_keys[sys_label], "") or "")
            print(f"[{call_n:4}/{total_calls}] {qid} System {sys_label} ...",
                  end=" ", flush=True)
            t0 = time.time()
            scores = _judge_one(client, q_text, expected, answer)
            elapsed = time.time() - t0
            print(f"total={scores['total']}/8  ({elapsed:.1f}s)")
            time.sleep(1)

            score_rows.append({
                "question_id":            qid,
                "domain":                 domain,
                "tier":                   tier,
                "system":                 sys_label,
                "factual_accuracy":       scores.get("factual_accuracy", 0),
                "completeness":           scores.get("completeness", 0),
                "confidence_calibration": scores.get("confidence_calibration", 0),
                "total":                  scores.get("total", 0),
                "rationale":              scores.get("rationale", ""),
                "question_text":          q_text,
                "expected_answer":        expected,
                "system_answer":          answer,
            })

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = RESULTS_DIR / f"benchmark_evaluation_{ts}.xlsx"
    _save_evaluation(score_rows, out_path)
    print(f"\nEvaluation saved: {out_path}")

    # Quick summary to stdout
    from collections import defaultdict
    by_sys: dict = defaultdict(list)
    for row in score_rows:
        by_sys[row["system"]].append(row["total"])
    print("\n── Quick Summary (Avg Total / 8) ──────────────────")
    for sys_label in sorted(by_sys):
        vals = by_sys[sys_label]
        avg  = sum(vals) / len(vals) if vals else 0
        print(f"  System {sys_label}: {avg:.2f}  (n={len(vals)})")
    print("────────────────────────────────────────────────────")


if __name__ == "__main__":
    main()
