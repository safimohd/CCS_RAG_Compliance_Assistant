"""
Benchmark runner — generates answers from all three systems for every question.

Systems
-------
A : Base LLM  (gpt-5-nano, no retrieval)
B : Existing RAG (ChromaDB + gpt-5-nano via rag_query.py)
C : Agentic RAG (query decomposition + multi-hop retrieval, agentic_rag.py)

Output
------
benchmark/results/benchmark_answers_{timestamp}.xlsx
  Columns: question_id | domain | tier | question_text | expected_answer
           system_a_answer | system_b_answer | system_c_answer
           system_c_sub_queries | system_c_steps

Usage
-----
# Dry-run (one question per domain):
    python run_benchmark.py --dry-run

# Full run, both domains:
    python run_benchmark.py

# Single domain:
    python run_benchmark.py --domain Healthcare
    python run_benchmark.py --domain RERA

# Tier filter:
    python run_benchmark.py --tier 1
    python run_benchmark.py --tier 1 --tier 2

# Limit per domain (for testing):
    python run_benchmark.py --limit 5

# Resume from a previous partial run:
    python run_benchmark.py --resume benchmark/results/benchmark_answers_20240101_120000.xlsx

OPENAI_API_KEY must be set in the environment.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

# ─── path setup ───────────────────────────────────────────────────────────────
HERE         = Path(__file__).resolve().parent          # benchmark/
APP_DIR      = HERE.parent                              # 05_App/
PROJECT_ROOT = APP_DIR.parent                           # project root
sys.path.insert(0, str(APP_DIR))

RESULTS_DIR = HERE / "results"
RESULTS_DIR.mkdir(exist_ok=True)

# ─── LLM config (mirrors rag_query.py) ────────────────────────────────────────
LLM_MODEL = os.getenv("LLM_MODEL", "gpt-5-nano")

# ─── System A prompt ──────────────────────────────────────────────────────────
SYSTEM_A_PROMPT = """\
You are an Indian regulatory compliance expert. Answer the question using your
general knowledge of Indian laws and regulations. Be specific, cite relevant
sections where known, and flag any uncertainty. Structure your answer as:
1. Direct answer
2. Relevant legal provision (Act / Section / Rule)
3. Key compliance takeaway
"""


# ─── System A ─────────────────────────────────────────────────────────────────
def system_a_answer(client, question_text: str) -> str:
    """Base LLM — no retrieval."""
    try:
        rsp = client.chat.completions.create(
            model=LLM_MODEL,
            messages=[
                {"role": "system", "content": SYSTEM_A_PROMPT},
                {"role": "user",   "content": question_text},
            ],
        )
        return rsp.choices[0].message.content.strip()
    except Exception as exc:
        return f"[ERROR: {exc}]"


# ─── System B ─────────────────────────────────────────────────────────────────
def system_b_answer(question_text: str, domain: str, store, client) -> str:
    """Existing RAG pipeline (mirrors rag_query.py query flow)."""
    try:
        from rag_query import query_rag
        result = query_rag(question_text, domain=domain, store=store, client=client)
        return result.get("answer", str(result))
    except ImportError:
        # Fallback: inline the retrieval + answer logic if rag_query API differs
        return _system_b_inline(question_text, domain, store, client)
    except Exception as exc:
        return f"[ERROR: {exc}]"


def _system_b_inline(question_text: str, domain: str, store, client) -> str:
    """Inline System B if rag_query.query_rag is unavailable."""
    try:
        TOP_K       = 6
        CANDIDATE_K = 24
        TIER_WEIGHT = {"A": 1.00, "B": 0.95, "C": 0.80, "D": 0.55, "E": 0.35}

        space = "l2"
        try:
            space = ((store._collection.metadata or {}).get("hnsw:space") or "l2").lower()
        except Exception:
            pass

        results = store.similarity_search_with_relevance_scores(
            question_text, k=CANDIDATE_K, filter={"domain": domain})

        scored = []
        for doc, raw in results:
            md = doc.metadata or {}
            if md.get("tier", "E") == "E":
                continue
            cos = raw if space.startswith("cosine") else 1.0 - (1.0 - raw) ** 2
            weighted = cos * TIER_WEIGHT.get(md.get("tier", "C"), 0.80)
            scored.append({
                "content": doc.page_content,
                "title":   md.get("title", "Unknown"),
                "tier":    md.get("tier", "?"),
                "weighted": weighted,
            })

        scored.sort(key=lambda x: x["weighted"], reverse=True)
        top = scored[:TOP_K]

        context = "\n\n".join(
            f"[{i}] Title: {p['title']} | Tier: {p['tier']}\n{p['content']}"
            for i, p in enumerate(top, 1)
        )

        SYSTEM_B_PROMPT = """\
You are an Indian regulatory compliance assistant. Answer ONLY from the
retrieved sources supplied. Structure every answer as:
1. A direct answer.
2. Sources relied on, cited by document TITLE.
3. The specific regulatory provision (section, rule, clause).
4. A concise compliance takeaway.
Hard rules: never invent citations; flag DRAFT material; if evidence is
insufficient, say so explicitly."""

        rsp = client.chat.completions.create(
            model=LLM_MODEL,
            messages=[
                {"role": "system", "content": SYSTEM_B_PROMPT},
                {"role": "user",
                 "content": f"Domain: {domain}\n\nRetrieved sources:\n{context}\n\nQuestion: {question_text}"},
            ],
        )
        return rsp.choices[0].message.content.strip()
    except Exception as exc:
        return f"[ERROR: {exc}]"


# ─── System C ─────────────────────────────────────────────────────────────────
def system_c_answer(question_text: str, domain: str, store, client) -> dict:
    """Agentic RAG."""
    try:
        from agentic_rag import agentic_answer
        return agentic_answer(question_text, domain, store=store, client=client)
    except Exception as exc:
        return {"answer": f"[ERROR: {exc}]", "sub_queries": [], "steps": []}


# ─── helpers ──────────────────────────────────────────────────────────────────
def _load_chromadb():
    from langchain_huggingface import HuggingFaceEmbeddings
    from langchain_chroma import Chroma
    index_dir = APP_DIR / "chroma_db"
    if not index_dir.exists():
        sys.exit(f"ERROR: ChromaDB index not found at {index_dir}. "
                 "Run ingest_corpus.py --rebuild first.")
    emb = HuggingFaceEmbeddings(
        model_name="BAAI/bge-small-en-v1.5",
        encode_kwargs={"normalize_embeddings": True},
    )
    return Chroma(
        collection_name="regulatory_corpus",
        embedding_function=emb,
        persist_directory=str(index_dir),
    )


def _load_questions(domain_filter: str | None, tier_filter: list[int],
                    limit: int | None) -> list[dict]:
    from questions_healthcare import HEALTHCARE_QUESTIONS
    from questions_rera import RERA_QUESTIONS

    questions = []
    for q in HEALTHCARE_QUESTIONS:
        questions.append({**q, "domain": "Healthcare"})
    for q in RERA_QUESTIONS:
        questions.append({**q, "domain": "RERA"})

    if domain_filter:
        questions = [q for q in questions if q["domain"] == domain_filter]
    if tier_filter:
        questions = [q for q in questions if q["tier"] in tier_filter]
    if limit:
        # Take `limit` per domain
        from collections import defaultdict
        by_domain: dict = defaultdict(list)
        for q in questions:
            by_domain[q["domain"]].append(q)
        questions = []
        for qs in by_domain.values():
            questions.extend(qs[:limit])

    return questions


def _save_results(rows: list[dict], path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    cols = [
        ("question_id",          14),
        ("domain",               12),
        ("tier",                  6),
        ("question_text",        55),
        ("expected_answer",      70),
        ("system_a_answer",      70),
        ("system_b_answer",      70),
        ("system_c_answer",      70),
        ("system_c_sub_queries", 40),
        ("system_c_steps",       40),
    ]

    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ci, (col_key, col_w) in enumerate(cols, 1):
        cell = ws.cell(1, ci, col_key.replace("_", " ").title())
        cell.font      = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = border
        ws.column_dimensions[get_column_letter(ci)].width = col_w
    ws.row_dimensions[1].height = 25

    domain_fills = {
        "Healthcare": PatternFill("solid", fgColor="E2EFDA"),
        "RERA":       PatternFill("solid", fgColor="DDEEFF"),
    }

    for ri, row in enumerate(rows, 2):
        fill = domain_fills.get(row.get("domain", ""), PatternFill())
        for ci, (col_key, _) in enumerate(cols, 1):
            val = row.get(col_key, "")
            if isinstance(val, (list, dict)):
                val = json.dumps(val, ensure_ascii=False)
            cell = ws.cell(ri, ci, val)
            cell.font      = Font(name="Calibri", size=9)
            cell.fill      = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border    = border
        ws.row_dimensions[ri].height = 80

    ws.freeze_panes = "A2"
    wb.save(str(path))


def _load_existing(path: Path) -> dict:
    """Load already-completed question_ids from a previous run."""
    import openpyxl
    done: dict[str, dict] = {}
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        qid = rec.get("Question Id") or rec.get("question_id")
        if qid:
            done[str(qid)] = rec
    return done


# ─── main ─────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="CCS Benchmark Runner")
    parser.add_argument("--domain",   choices=["Healthcare", "RERA"],
                        help="Run for one domain only")
    parser.add_argument("--tier",     type=int, action="append", dest="tiers",
                        help="Filter to tier(s) (repeatable)")
    parser.add_argument("--limit",    type=int,
                        help="Max questions per domain (for testing)")
    parser.add_argument("--dry-run",  action="store_true",
                        help="Run 1 question per domain, print results, no save")
    parser.add_argument("--resume",   type=Path,
                        help="Resume from a partial results file")
    parser.add_argument("--no-system-a", action="store_true",
                        help="Skip System A (base LLM)")
    parser.add_argument("--no-system-b", action="store_true",
                        help="Skip System B (existing RAG)")
    parser.add_argument("--no-system-c", action="store_true",
                        help="Skip System C (agentic RAG)")
    args = parser.parse_args()

    # Validate API key
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        sys.exit("ERROR: OPENAI_API_KEY not set in environment.")

    from openai import OpenAI
    client = OpenAI(api_key=api_key)

    # Load questions
    questions = _load_questions(
        domain_filter=args.domain,
        tier_filter=args.tiers or [],
        limit=1 if args.dry_run else args.limit,
    )
    print(f"Loaded {len(questions)} question(s).")

    # Load ChromaDB (needed for B and C)
    need_store = not (args.no_system_b and args.no_system_c)
    store = _load_chromadb() if need_store else None
    if store:
        print(f"ChromaDB loaded: {APP_DIR / 'chroma_db'}")

    # Load previous results if resuming
    done_map: dict = {}
    if args.resume and args.resume.exists():
        done_map = _load_existing(args.resume)
        print(f"Resuming — {len(done_map)} question(s) already done.")

    # Output file
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = RESULTS_DIR / f"benchmark_answers_{ts}.xlsx"

    rows: list[dict] = list(done_map.values())  # carry over prior results

    total = len(questions)
    for idx, q in enumerate(questions, 1):
        qid = q["question_id"]
        if qid in done_map:
            print(f"[{idx:3}/{total}] {qid} — SKIP (already done)")
            continue

        domain = q["domain"]
        text   = q["question_text"]
        print(f"[{idx:3}/{total}] {qid} ({domain} T{q['tier']}) ...", end=" ", flush=True)

        row: dict[str, Any] = {
            "question_id":     qid,
            "domain":          domain,
            "tier":            q["tier"],
            "question_text":   text,
            "expected_answer": q.get("expected_answer", ""),
            "system_a_answer": "",
            "system_b_answer": "",
            "system_c_answer": "",
            "system_c_sub_queries": "",
            "system_c_steps":  "",
        }

        t0 = time.time()

        if not args.no_system_a:
            row["system_a_answer"] = system_a_answer(client, text)
            print("A✓", end=" ", flush=True)

        if not args.no_system_b and store:
            row["system_b_answer"] = system_b_answer(text, domain, store, client)
            print("B✓", end=" ", flush=True)

        if not args.no_system_c and store:
            c_result = system_c_answer(text, domain, store, client)
            row["system_c_answer"]      = c_result.get("answer", "")
            row["system_c_sub_queries"] = c_result.get("sub_queries", [])
            row["system_c_steps"]       = c_result.get("steps", [])
            print("C✓", end=" ", flush=True)

        elapsed = time.time() - t0
        print(f"({elapsed:.1f}s)")

        rows.append(row)

        if args.dry_run:
            print("\n--- DRY RUN RESULTS ---")
            for sys_key in ("system_a_answer", "system_b_answer", "system_c_answer"):
                val = row.get(sys_key, "")
                if val:
                    print(f"\n[{sys_key.upper()}]\n{val[:500]}...")
            print("\n--- END DRY RUN ---")

    if not args.dry_run and rows:
        _save_results(rows, out_path)
        print(f"\nResults saved: {out_path}")
        print(f"Total rows: {len(rows)}")
    elif args.dry_run:
        print("\nDry run complete. No file saved.")


if __name__ == "__main__":
    main()
