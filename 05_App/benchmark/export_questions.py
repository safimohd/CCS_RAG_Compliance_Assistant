"""
Export benchmark questions to Excel for submission.

Produces:
  benchmark/healthcare_benchmark.xlsx
  benchmark/rera_benchmark.xlsx

Columns: question_id | tier | question_text | expected_answer
         source_citation | source_url | known_trap

Run from any directory:
    python 05_App/benchmark/export_questions.py
"""

from pathlib import Path
import sys

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("question_id",     14),
    ("tier",             6),
    ("question_text",   60),
    ("expected_answer", 80),
    ("source_citation", 40),
    ("source_url",      30),
    ("known_trap",      40),
]

HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")   # dark navy
TIER_FILLS = {
    1: PatternFill("solid", fgColor="DEEAF1"),   # lightest blue
    2: PatternFill("solid", fgColor="BDD7EE"),
    3: PatternFill("solid", fgColor="9DC3E6"),
    4: PatternFill("solid", fgColor="5BA3C9"),
    5: PatternFill("solid", fgColor="2E75B6"),   # deepest
}
TIER_FONT_DARK = {4, 5}

thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _write_sheet(wb, questions: list[dict], title: str) -> None:
    ws = wb.create_sheet(title=title)

    # Header
    for col_idx, (col_key, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_key.replace("_", " ").title())
        cell.font      = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 30

    # Data rows
    for row_idx, q in enumerate(questions, 2):
        tier = q.get("tier", 1)
        fill = TIER_FILLS.get(tier, TIER_FILLS[1])
        font_color = "FFFFFF" if tier in TIER_FONT_DARK else "000000"
        for col_idx, (col_key, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=q.get(col_key, ""))
            cell.font      = Font(name="Calibri", size=10, color=font_color)
            cell.fill      = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border    = BORDER
        ws.row_dimensions[row_idx].height = 60

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"


def build_workbook(questions: list[dict], domain_label: str) -> openpyxl.Workbook:
    from collections import defaultdict
    wb = openpyxl.Workbook()
    wb.remove(wb.active)            # remove default sheet

    # Summary sheet
    ws_sum = wb.create_sheet("Summary", 0)
    ws_sum["A1"] = f"CCS Benchmark — {domain_label}"
    ws_sum["A1"].font = Font(bold=True, size=14, name="Calibri")
    ws_sum["A2"] = f"Total questions: {len(questions)}"
    tier_names = {
        1: "Factual Recall",
        2: "Conceptual Understanding",
        3: "Scenario Application",
        4: "Multi-step Reasoning",
        5: "Adversarial / Trap",
    }
    by_tier = defaultdict(list)
    for q in questions:
        by_tier[q["tier"]].append(q)
    row = 4
    ws_sum.cell(row, 1, "Tier").font = Font(bold=True)
    ws_sum.cell(row, 2, "Name").font = Font(bold=True)
    ws_sum.cell(row, 3, "Count").font = Font(bold=True)
    row += 1
    for tier in sorted(by_tier):
        ws_sum.cell(row, 1, f"T{tier}")
        ws_sum.cell(row, 2, tier_names.get(tier, ""))
        ws_sum.cell(row, 3, len(by_tier[tier]))
        row += 1
    ws_sum.column_dimensions["A"].width = 8
    ws_sum.column_dimensions["B"].width = 32
    ws_sum.column_dimensions["C"].width = 10

    # All questions (one sheet)
    _write_sheet(wb, questions, "All Questions")

    # Per-tier sheets
    for tier in sorted(by_tier):
        safe_name = tier_names.get(tier, "").replace("/", "-")
        _write_sheet(wb, by_tier[tier], f"T{tier} {safe_name}")

    return wb


def main():
    from questions_healthcare import HEALTHCARE_QUESTIONS
    from questions_rera import RERA_QUESTIONS

    out_dir = HERE

    # Healthcare
    hc_path = out_dir / "healthcare_benchmark.xlsx"
    wb_hc = build_workbook(HEALTHCARE_QUESTIONS, "Healthcare")
    wb_hc.save(str(hc_path))
    print(f"Saved: {hc_path}  ({len(HEALTHCARE_QUESTIONS)} questions)")

    # RERA
    re_path = out_dir / "rera_benchmark.xlsx"
    wb_re = build_workbook(RERA_QUESTIONS, "RERA")
    wb_re.save(str(re_path))
    print(f"Saved: {re_path}  ({len(RERA_QUESTIONS)} questions)")

    print("Done.")


if __name__ == "__main__":
    main()
